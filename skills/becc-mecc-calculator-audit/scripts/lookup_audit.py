#!/usr/bin/env python3
"""Lookup, material-mapping and emission-factor-table audit (stages 5, 11, 24).

For every lookup-type call (VLOOKUP, HLOOKUP, LOOKUP, MATCH, XLOOKUP, XMATCH,
INDEX, SUMIF(S), COUNTIF(S), INDIRECT/OFFSET) it records:
  * the table / key range and the sheet it lives on
  * match mode — approximate matching on unsorted text keys silently returns
    the wrong row (VLOOKUP 4th arg omitted/TRUE, MATCH 3rd arg omitted/1, LOOKUP)
  * the column returned and that column's header (is this really the A1-A3 factor?)
  * TRUNCATED: keys exist on the target sheet beyond the end of the range
  * DUPLICATE_KEYS in the key column (ambiguous lookups — first match wins)
  * INCONSISTENT_RANGE: the same table addressed with extents that stop short of the data
    (RANGE_EXTENT_VARIES = differing headroom beyond the data: benign but untidy)
  * MISALIGNED_RETURN_RANGE for XLOOKUP
  * fallback chains, e.g. IFERROR(lookup(country&material), lookup("Global"&material))
Then a material-mapping matrix (block label → table sheet) with a keyword
heuristic flagging rows whose material label does not match the table read,
and a profile of each lookup target table (records, source/year/version/
geography/unit/boundary columns, negatives, zeros, outliers).

Outputs: lookups.csv, lookup_audit.md, material_mapping.csv, ef_tables.csv
Usage: python lookup_audit.py WORKBOOK [--out DIR]
"""
from __future__ import annotations

import argparse
import os
import re
import statistics
from collections import Counter, defaultdict

from common import (load_cells, index_cells, parse_calls, split_sheet, parse_address, area_bounds,
                    row_label, write_csv, ensure_dir, short)
from excel_compat import base_name
from openpyxl.utils import get_column_letter

MATERIALS = ["concrete", "cement", "admixture", "aggregate", "steel", "rebar", "reinforcement", "aluminium",
             "aluminum", "glass", "timber", "wood", "tile", "paint", "brick", "waterproof", "carpet", "copper",
             "mineral", "wool", "insulation", "paper", "plaster", "gypsum", "plastic", "sealant", "precast",
             "ggbs", "fly", "water", "stainless", "section", "diesel", "petrol", "electric"]
META_HINTS = {"source": ["source", "reference", "database", "dataset", "epd", "publisher", "link", "url"],
              "year": ["year", "date", "valid", "published", "expiry"],
              "version": ["version", "ver.", "edition"],
              "geography": ["country", "region", "geograph", "location", "origin"],
              "unit": ["unit", "declared"],
              "boundary": ["a1", "a2", "a3", "a4", "a5", "module", "stage", "boundary", "c1", "c2", "c3", "c4"]}


def material_keywords(text: str) -> set:
    t = (text or "").lower().replace("_", " ")
    ks = {m for m in MATERIALS if m in t}
    if "aluminum" in ks:
        ks.add("aluminium")
    if "rebar" in ks or "reinforcement" in ks:
        ks.add("steel")
    return ks


def block_label(sidx, row, cols=(1, 2), max_up=60):
    parts = []
    for col in cols:
        for r in range(row, max(0, row - max_up), -1):
            c = sidx.get((r, col))
            if c is not None and c.kind == "text" and str(c.value).strip():
                parts.append(str(c.value).strip()); break
    return " | ".join(parts)


def resolve_range(arg: str, host_sheet: str):
    arg = arg.strip()
    if re.match(r"^('?[^()]+'?!)?\$?[A-Za-z]{1,3}\$?\d*(:\$?[A-Za-z]{1,3}\$?\d*)?$", arg):
        book, sh, addr = split_sheet(arg)
        kind, data = parse_address(addr)
        if kind in ("cell", "area", "cols", "rows"):
            return {"book": book, "sheet": sh or host_sheet, "bounds": area_bounds(kind, data), "kind": kind, "text": arg}
    return None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--password")
    ap.add_argument("--out", default="audit_out")
    a = ap.parse_args()
    out = ensure_dir(a.out)
    wb_f, wb_v, real, cells = load_cells(a.workbook, a.password)
    idx = index_cells(cells)

    def last_row_in_col(sheet, col):
        rs = [r for (r, c), cell in idx.get(sheet, {}).items() if c == col and cell.value not in (None, "")]
        return max(rs) if rs else 0

    def key_values(sheet, col, r1, r2):
        return sorted((r, cell.value) for (r, c), cell in idx.get(sheet, {}).items() if c == col and r1 <= r <= r2)

    lookups = defaultdict(lambda: {"hosts": [], "host_sheets": Counter(), "labels": Counter()})
    mapping = defaultdict(Counter)
    mapping_rows = defaultdict(set)
    table_usage = defaultdict(set)
    fallback_patterns = Counter()

    for c in cells:
        if not c.is_formula:
            continue
        calls = parse_calls(c.formula)
        if not calls:
            continue
        up = c.formula.upper()
        if "IFERROR(" in up and up.count("LOOKUP(") >= 2:
            m = re.findall(r'CONCAT\("([^"]+)"', c.formula)
            if m:
                fallback_patterns[(c.sheet, "fallback key prefix " + ", ".join(sorted(set(m))))] += 1
        sidx = idx[c.sheet]
        label = block_label(sidx, c.row)
        for call in calls:
            fn = base_name(call.name)
            args = call.args
            rec = None
            if fn in ("VLOOKUP", "HLOOKUP") and len(args) >= 3:
                mode = "exact" if len(args) >= 4 and args[3].upper() in ("FALSE", "0") else "APPROXIMATE"
                rec = (fn, args[1], args[2], mode, resolve_range(args[1], c.sheet))
            elif fn == "MATCH" and len(args) >= 2:
                mode = "exact" if len(args) >= 3 and args[2] == "0" else ("APPROXIMATE" if len(args) < 3 or args[2] in ("1", "-1") else args[2])
                rec = (fn, args[1], "", mode, resolve_range(args[1], c.sheet))
            elif fn == "XLOOKUP" and len(args) >= 3:
                mm = args[4] if len(args) >= 5 and args[4] else "0"
                sm = args[5] if len(args) >= 6 and args[5] else "1"
                mode = {"0": "exact", "-1": "exact-or-next-smaller", "1": "exact-or-next-larger", "2": "wildcard"}.get(mm, mm)
                mode += "" if sm == "1" else f", search_mode {sm}"
                mode += "" if len(args) >= 4 and args[3] else ", no if_not_found"
                rec = (fn, args[1], args[2], mode, resolve_range(args[1], c.sheet))
            elif fn == "XMATCH" and len(args) >= 2:
                rec = (fn, args[1], "", "exact" if len(args) < 3 or args[2] == "0" else f"match_mode {args[2]}",
                       resolve_range(args[1], c.sheet))
            elif fn == "LOOKUP" and len(args) >= 2:
                rec = (fn, args[1], "", "APPROXIMATE (LOOKUP is always approximate)", resolve_range(args[1], c.sheet))
            elif fn in ("SUMIF", "COUNTIF", "AVERAGEIF") and len(args) >= 2:
                rec = (fn, args[0], args[2] if len(args) > 2 else "", "criteria", resolve_range(args[0], c.sheet))
            elif fn in ("SUMIFS", "COUNTIFS", "AVERAGEIFS", "MAXIFS", "MINIFS") and len(args) >= 2:
                rng = args[1] if fn != "COUNTIFS" else args[0]
                rec = (fn, rng, args[0] if fn != "COUNTIFS" else "", "criteria", resolve_range(rng, c.sheet))
            elif fn == "INDEX" and len(args) >= 2:
                rec = (fn, args[0], ",".join(args[1:3]), "positional", resolve_range(args[0], c.sheet))
            elif fn in ("INDIRECT", "OFFSET"):
                rec = (fn, args[0] if args else "", "", "DYNAMIC (not statically traceable)", None)
            if rec is None:
                continue
            fn_, rng, colidx, mode, tbl = rec
            L = lookups[(fn_, rng.replace("$", ""), colidx, mode)]
            L["hosts"].append(c)
            L["host_sheets"][c.sheet] += 1
            if label:
                L["labels"][label] += 1
            L["tbl"] = tbl
            if tbl:
                mapping[(c.sheet, label)][tbl["sheet"]] += 1
                if len(mapping_rows[(c.sheet, label)]) < 8:
                    mapping_rows[(c.sheet, label)].add(c.row)
                if tbl["bounds"] and fn_ != "HLOOKUP":
                    r1, c1, r2, c2 = tbl["bounds"]
                    table_usage[(tbl["sheet"], c1)].add((r1, r2))

    rows = []
    for (fn, rng, colidx, mode), L in lookups.items():
        tbl = L.get("tbl")
        issues, detail, ret_header = set(), [], ""
        if "APPROX" in mode:
            issues.add("APPROXIMATE_MATCH")
        if "DYNAMIC" in mode:
            issues.add("DYNAMIC_REFERENCE")
        if tbl and tbl["book"]:
            issues.add("EXTERNAL_TABLE")
        if tbl and tbl["bounds"] and not tbl["book"]:
            r1, c1, r2, c2 = tbl["bounds"]
            sh = tbl["sheet"]
            if sh not in wb_f.sheetnames:
                issues.add("MISSING_SHEET")
            elif fn != "HLOOKUP":
                last = last_row_in_col(sh, c1)
                if r2 < 1048576 and last > r2:
                    issues.add("TRUNCATED")
                    detail.append(f"key column {get_column_letter(c1)} has data to row {last} but range ends at {r2}")
                kv = key_values(sh, c1, r1, min(r2, max(last, r1)))
                nonblank = [(r, v) for r, v in kv if v not in (None, "") and not (isinstance(v, str) and v.startswith("#"))]
                norm = Counter(str(v).strip().lower() for r, v in nonblank)
                dups = {k: n for k, n in norm.items() if n > 1}
                if dups and fn in ("VLOOKUP", "MATCH", "XLOOKUP", "XMATCH", "LOOKUP", "SUMIF", "SUMIFS"):
                    issues.add("DUPLICATE_KEYS")
                    detail.append("duplicate keys: " + "; ".join(f"{short(k, 40)}×{n}" for k, n in list(dups.items())[:6]))
                padded = [str(v) for r, v in nonblank if isinstance(v, str) and v != v.strip()]
                if padded:
                    detail.append(f"{len(padded)} key(s) with leading/trailing spaces, e.g. {padded[0]!r}")
                if fn == "VLOOKUP" and colidx.isdigit():
                    ci = int(colidx)
                    if ci > (c2 - c1 + 1):
                        issues.add("COLINDEX_OUT_OF_RANGE")
                    hc = idx.get(sh, {}).get((1, c1 + ci - 1))
                    ret_header = "" if hc is None or str(hc.value).startswith("=") else str(hc.value)
                elif fn == "XLOOKUP":
                    rr = resolve_range(colidx, L["hosts"][0].sheet)
                    if rr and rr["bounds"]:
                        hc = idx.get(rr["sheet"], {}).get((1, rr["bounds"][1]))
                        ret_header = str(hc.value) if hc is not None else ""
                        if rr["bounds"][0] != r1 or rr["bounds"][2] != r2:
                            issues.add("MISALIGNED_RETURN_RANGE")
                            detail.append(f"lookup rows {r1}-{r2} vs return rows {rr['bounds'][0]}-{rr['bounds'][2]}")
                if "APPROXIMATE_MATCH" in issues and nonblank:
                    vals = [v for _, v in nonblank]
                    try:
                        sortd = all(vals[i] <= vals[i + 1] for i in range(len(vals) - 1))
                    except TypeError:
                        sortd = False
                    detail.append("keys sorted ascending" if sortd else "keys NOT sorted ascending → approximate match can return wrong row")
        hosts = L["hosts"]
        rows.append({"function": fn, "range": rng, "col_or_return": colidx, "match_mode": mode,
                     "target_sheet": tbl["sheet"] if tbl else "", "returned_header": short(ret_header, 60),
                     "issues": ", ".join(sorted(issues)), "detail": "; ".join(detail), "uses": len(hosts),
                     "host_sheets": ", ".join(f"{k}×{v}" for k, v in L["host_sheets"].most_common(4)),
                     "example_host": hosts[0].full, "example_formula": short(hosts[0].formula, 220),
                     "row_labels": "; ".join(short(k, 40) for k, _ in L["labels"].most_common(4))})

    incons = []
    for (sh, kc), exts in table_usage.items():
        if len(exts) > 1:
            last = last_row_in_col(sh, kc)
            incons.append({"target": f"{sh}!{get_column_letter(kc)}", "sheet": sh, "extents": sorted(exts),
                           "data_last_row": last, "short": any(r2 < last for r1, r2 in exts)})
    for r in rows:
        for inc in incons:
            if r["target_sheet"] == inc["sheet"]:
                tag = "INCONSISTENT_RANGE" if inc["short"] else "RANGE_EXTENT_VARIES"
                if tag not in r["issues"]:
                    r["issues"] = ", ".join(x for x in [r["issues"], tag] if x)

    map_rows = []
    for (hs, label), cnt in mapping.items():
        lk = material_keywords(label)
        for tsheet, n in cnt.items():
            tk = material_keywords(tsheet)
            susp = bool(lk and tk and not (lk & tk)) and not tsheet.lower().startswith(("vehicle", "maritime", "dropdown", "construction"))
            map_rows.append({"host_sheet": hs, "block_label": short(label, 80), "target_sheet": tsheet, "lookups": n,
                             "label_keywords": ",".join(sorted(lk)), "target_keywords": ",".join(sorted(tk)),
                             "flag": "CHECK: label/target material mismatch" if susp else "",
                             "sample_rows": ",".join(str(x) for x in sorted(mapping_rows[(hs, label)]))})
    map_rows.sort(key=lambda r: (r["flag"] == "", r["host_sheet"], r["block_label"]))

    host_load = Counter()
    for r in rows:
        for part in r["host_sheets"].split(", "):
            if "×" in part:
                h, n = part.rsplit("×", 1); host_load[h] += int(n)
    targets = sorted({r["target_sheet"] for r in rows if r["target_sheet"] in idx and host_load[r["target_sheet"]] < 50})
    prof = []
    for sh in targets:
        sidx = idx[sh]
        hdr = {c: str(cell.value).strip() for (r, c), cell in sidx.items() if r == 1 and cell.value is not None}
        nrec = len({r for (r, c) in sidx if r > 1})
        meta = {k: [h for h in hdr.values() if any(x in h.lower() for x in v)] for k, v in META_HINTS.items()}
        numeric_issues = []
        for c, h in hdr.items():
            vals = [cell.value for (r, cc), cell in sidx.items() if cc == c and r > 1]
            nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            if len(nums) < 2:
                continue
            neg = sum(1 for v in nums if v < 0)
            zer = sum(1 for v in nums if v == 0)
            med = statistics.median([abs(v) for v in nums]) or 0
            outl = [v for v in nums if med and abs(v) > 50 * med]
            if neg or zer or outl:
                numeric_issues.append(f"{h[:30]}: neg={neg} zero={zer} outliers(>50×median)={len(outl)}")
        prof.append({"sheet": sh, "records": nrec, "headers": short(" | ".join(hdr.values()), 400),
                     "has_source": bool(meta["source"]), "has_year": bool(meta["year"]),
                     "has_version": bool(meta["version"]), "has_geography": bool(meta["geography"]),
                     "has_unit": bool(meta["unit"]), "has_boundary": bool(meta["boundary"]),
                     "matched_meta_headers": short("; ".join(f"{k}: {', '.join(v)[:80]}" for k, v in meta.items() if v), 500),
                     "numeric_flags": "; ".join(numeric_issues)[:500]})

    rows.sort(key=lambda r: (r["issues"] in ("", "RANGE_EXTENT_VARIES"), r["target_sheet"], -r["uses"]))
    write_csv(os.path.join(out, "lookups.csv"), rows)
    write_csv(os.path.join(out, "material_mapping.csv"), map_rows)
    write_csv(os.path.join(out, "ef_tables.csv"), prof)

    Lm = [f"# Lookup & mapping audit — {os.path.basename(a.workbook)}\n",
          f"{len(rows)} distinct lookup expressions across {sum(r['uses'] for r in rows)} call sites.\n"]
    ic = Counter(i.strip() for r in rows for i in r["issues"].split(",") if i.strip())
    Lm.append("Issue counts (distinct expressions): " + ", ".join(f"{k}: {v}" for k, v in ic.most_common()))
    Lm.append("\nHeader detection for metadata columns is keyword-based — confirm by opening the table "
              "(e.g. a 'Type' column may hold 'Country'/'Global' rather than a source).")
    if incons:
        Lm.append("\n## Same table, different extents\n")
        for inc in sorted(incons, key=lambda x: not x["short"]):
            Lm.append(f"- `{inc['target']}` addressed with row extents {inc['extents']} (data to row {inc['data_last_row']})"
                      + (" — **at least one extent stops short of the data**" if inc["short"] else " — benign headroom difference"))
    if fallback_patterns:
        Lm.append("\n## Fallback chains (IFERROR to a second lookup)\n")
        for (sh, p), n in fallback_patterns.most_common(20):
            Lm.append(f"- {sh}: {p} ×{n} — a missing specific factor silently falls back to a generic one; "
                      "confirm this is documented methodology and visible to users")
    Lm.append("\n## Lookups with issues\n")
    for r in [r for r in rows if r["issues"] and r["issues"] != "RANGE_EXTENT_VARIES"][:150]:
        Lm.append(f"- **{r['function']}** `{r['range']}` col/ret `{r['col_or_return']}` ({r['match_mode']}) → "
                  f"{r['target_sheet']} [{r['returned_header']}] — **{r['issues']}** {r['detail']}  \n"
                  f"  uses {r['uses']} ({r['host_sheets']}), e.g. {r['example_host']}: `{r['example_formula']}`")
    Lm.append("\n## All lookups by target (returned column header)\n")
    Lm.append("| Function | Range | Col/return | Header returned | Mode | Uses | Labels |\n|---|---|---|---|---|---|---|")
    for r in rows[:200]:
        Lm.append(f"| {r['function']} | `{r['range']}` | {r['col_or_return']} | {r['returned_header']} | {r['match_mode']} | "
                  f"{r['uses']} | {r['row_labels']} |")
    Lm.append("\n## Material mapping (block label → table sheet)\n")
    Lm.append("| Host sheet | Block label | Target sheet | Lookups | Flag | Sample rows |\n|---|---|---|---|---|---|")
    for m in map_rows[:250]:
        Lm.append(f"| {m['host_sheet']} | {m['block_label']} | {m['target_sheet']} | {m['lookups']} | {m['flag']} | {m['sample_rows']} |")
    Lm.append("\n## Emission-factor / reference table profiles\n")
    Lm.append("| Sheet | Records | Source | Year | Version | Geography | Unit | LC boundary | Numeric flags |\n|---|---|---|---|---|---|---|---|---|")
    yn = lambda b: "✓" if b else "✗"
    for p in prof:
        Lm.append(f"| {p['sheet']} | {p['records']} | {yn(p['has_source'])} | {yn(p['has_year'])} | {yn(p['has_version'])} | "
                  f"{yn(p['has_geography'])} | {yn(p['has_unit'])} | {yn(p['has_boundary'])} | {p['numeric_flags']} |")
    with open(os.path.join(out, "lookup_audit.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(Lm) + "\n")
    print(f"lookup audit → {out}/lookup_audit.md ; issues: {dict(ic)}")


if __name__ == "__main__":
    main()
