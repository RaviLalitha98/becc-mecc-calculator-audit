#!/usr/bin/env python3
"""Validate findings.json and build the findings register.

Outputs in --out:
  findings_register.xlsx  formatted, filterable register + Summary sheet (severity × class, category counts)
  findings_register.md    Markdown table for the report
  findings_counts.json    counts for the executive summary

Schema and enumerations: references/severity-confidence-findings.md
Usage: python findings_register.py findings.json --out audit_out
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter

SEVERITIES = ["Critical", "High", "Medium", "Low", "Observation"]
CONFIDENCES = ["Confirmed", "High", "Medium", "Low", "Requires methodology confirmation"]
CLASSES = ["Confirmed error", "Probable error requiring review", "Unusual but potentially intentional logic",
           "Methodology/design issue", "Maintainability risk", "Version-control / workbook-evolution risk"]
REQUIRED = ["id", "title", "class", "severity", "confidence", "sheet", "range", "category",
            "current_logic", "issue", "impact", "recommendation"]
COLUMNS = ["id", "severity", "confidence", "class", "category", "lifecycle_module", "material", "workbook", "sheet",
           "range", "title", "current_logic", "expected_logic", "linked_sources", "evidence", "issue", "impact",
           "recommendation", "proposed_formula", "verification", "status"]
FILL = {"Critical": "C00000", "High": "ED7D31", "Medium": "FFC000", "Low": "A9D08E", "Observation": "BDD7EE"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("findings")
    ap.add_argument("--out", default="audit_out")
    a = ap.parse_args()
    data = json.load(open(a.findings, encoding="utf-8"))
    errs = []
    for i, f in enumerate(data):
        for k in REQUIRED:
            if not str(f.get(k, "")).strip():
                errs.append(f"#{i} {f.get('id')}: missing {k}")
        if f.get("severity") not in SEVERITIES:
            errs.append(f"{f.get('id')}: severity {f.get('severity')!r} not in {SEVERITIES}")
        if f.get("confidence") not in CONFIDENCES:
            errs.append(f"{f.get('id')}: confidence {f.get('confidence')!r} not in {CONFIDENCES}")
        if f.get("class") not in CLASSES:
            errs.append(f"{f.get('id')}: class {f.get('class')!r} not in {CLASSES}")
        if f.get("class") == "Confirmed error" and f.get("confidence") in ("Low", "Requires methodology confirmation"):
            errs.append(f"{f.get('id')}: 'Confirmed error' with confidence {f.get('confidence')} is contradictory")
    for k, n in Counter(f.get("id") for f in data).items():
        if n > 1:
            errs.append(f"duplicate id {k}")
    if errs:
        print("VALIDATION ERRORS:\n  " + "\n  ".join(errs))
        sys.exit(1)
    data.sort(key=lambda f: (SEVERITIES.index(f["severity"]), CONFIDENCES.index(f["confidence"]), f["id"]))
    os.makedirs(a.out, exist_ok=True)

    sev = Counter(f["severity"] for f in data)
    cls = Counter(f["class"] for f in data)
    cat = Counter(f["category"] for f in data)
    matrix = Counter((f["severity"], f["class"]) for f in data)
    counts = {"total": len(data), "by_severity": {s: sev.get(s, 0) for s in SEVERITIES},
              "by_class": {c: cls.get(c, 0) for c in CLASSES}, "by_category": dict(cat.most_common()),
              "by_confidence": dict(Counter(f["confidence"] for f in data))}
    json.dump(counts, open(os.path.join(a.out, "findings_counts.json"), "w"), indent=2)

    esc = lambda s: str(s or "").replace("|", "\\|").replace("\n", " ")
    L = ["| ID | Severity | Confidence | Class | Sheet | Cell/Range | Category | Current logic | Expected/reference logic | Issue | Potential impact | Recommendation |",
         "|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for f in data:
        L.append("| " + " | ".join(esc(x) for x in [f["id"], f["severity"], f["confidence"], f["class"], f["sheet"], f["range"],
                                                  f["category"], f"`{f['current_logic']}`", f.get("expected_logic", ""),
                                                  f["issue"], f["impact"], f["recommendation"]]) + " |")
    with open(os.path.join(a.out, "findings_register.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active; ws.title = "Findings"
    ws.append([c.replace("_", " ").title() for c in COLUMNS])
    for f in data:
        ws.append([str(f.get(c, "")) if isinstance(f.get(c, ""), (list, dict)) else f.get(c, "") for c in COLUMNS])
    widths = {"id": 8, "severity": 11, "confidence": 14, "class": 24, "category": 16, "title": 40, "sheet": 18,
              "range": 14, "current_logic": 45, "expected_logic": 45, "issue": 50, "impact": 50, "recommendation": 45,
              "evidence": 50, "linked_sources": 30, "proposed_formula": 40, "verification": 35}
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 16)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    si = COLUMNS.index("severity")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = "'" + cell.value  # store formulas as text, never as live formulas
        s = row[si].value
        if s in FILL:
            row[si].fill = PatternFill("solid", fgColor=FILL[s])
            row[si].font = Font(bold=True, color="FFFFFF" if s in ("Critical", "High") else "000000")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    sm = wb.create_sheet("Summary")
    sm.append(["Severity \\ Class"] + CLASSES + ["Total"])
    for s in SEVERITIES:
        sm.append([s] + [matrix.get((s, c), 0) for c in CLASSES] + [sev.get(s, 0)])
    sm.append(["Total"] + [cls.get(c, 0) for c in CLASSES] + [len(data)])
    sm.append([])
    sm.append(["Category", "Count"])
    for k, v in cat.most_common():
        sm.append([k, v])
    for cell in sm[1]:
        cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True)
    sm.column_dimensions["A"].width = 22
    for i in range(2, len(CLASSES) + 3):
        sm.column_dimensions[get_column_letter(i)].width = 18
    wb.save(os.path.join(a.out, "findings_register.xlsx"))
    print(f"{len(data)} findings valid → findings_register.xlsx/.md; by severity {counts['by_severity']}")


if __name__ == "__main__":
    main()
