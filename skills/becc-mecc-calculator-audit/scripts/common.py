"""Shared helpers for the BECC/MECC workbook audit scripts.

Everything here is read-only with respect to the source workbook. Scripts that
need to modify a workbook (recalc.py) always work on a copy.

Key ideas
---------
* Every non-empty cell is loaded into a `Cell` record holding both the formula
  text (formula load) and the cached value Excel last saved (data_only load).
* Formulas are normalised to a relative R1C1 form so that "the same formula
  copied down a column" has identical text in every row. Pattern breaks are
  then visible as text differences.
* A small tokenizer-based parser extracts function calls and their arguments so
  lookups and aggregations can be analysed semantically rather than by regex.
"""
from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import warnings
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Any, Iterable

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Windows consoles default to cp1252, which cannot print the arrows/subscripts
# used in help text and summaries; fall back to replacement instead of crashing.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(errors="replace")

try:
    import openpyxl
    from openpyxl.formula.tokenizer import Tokenizer, Token
    from openpyxl.utils import get_column_letter, column_index_from_string
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl --break-system-packages")

try:  # DataTableFormula only exists in newer openpyxl versions
    from openpyxl.worksheet.formula import DataTableFormula
except Exception:  # pragma: no cover
    DataTableFormula = ()  # type: ignore

EXCEL_ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "#NULL!",
                "#SPILL!", "#CALC!", "#GETTING_DATA", "#FIELD!", "#BLOCKED!", "#CONNECT!",
                "#BUSY!", "#UNKNOWN!")

# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

def decrypt_if_needed(path: str, password: str | None, workdir: str | None = None) -> str:
    """Return a path openpyxl can open.

    If the file is an encrypted OOXML container (password-to-open), decrypt a
    *copy* with the password the authorised owner supplied. Never attempts to
    guess or crack passwords. Sheet/workbook *structure* protection does not
    block openpyxl from reading formulas, so it needs no special handling.
    """
    with open(path, "rb") as fh:
        head = fh.read(8)
    is_ole = head.startswith(b"\xD0\xCF\x11\xE0")
    if not is_ole or path.lower().endswith(".xls"):
        return path
    if not password:
        raise SystemExit(
            f"{path} looks encrypted (OLE container). Supply the owner's password with --password. "
            "Do not attempt to bypass the protection.")
    try:
        import msoffcrypto  # type: ignore
    except ImportError:
        raise SystemExit("Encrypted workbook: pip install msoffcrypto-tool --break-system-packages")
    workdir = workdir or tempfile.mkdtemp(prefix="becc_dec_")
    out = os.path.join(workdir, "decrypted_" + os.path.basename(path))
    with open(path, "rb") as fin, open(out, "wb") as fout:
        of = msoffcrypto.OfficeFile(fin)
        of.load_key(password=password)
        of.decrypt(fout)
    return out


def load_books(path: str, password: str | None = None):
    """Load workbook twice: formulas and cached values. Returns (wb_f, wb_v, real_path)."""
    if path.lower().endswith(".xls"):
        raise SystemExit("Legacy .xls is not readable by openpyxl. Convert a COPY with "
                         "`soffice --headless --convert-to xlsx` and note in the report that the "
                         "audit used a converted copy (conversion can alter some formulas).")
    real = decrypt_if_needed(path, password)
    keep_vba = real.lower().endswith(".xlsm")
    wb_f = openpyxl.load_workbook(real, data_only=False, keep_vba=keep_vba, keep_links=True)
    wb_v = openpyxl.load_workbook(real, data_only=True, keep_links=True)
    return wb_f, wb_v, real


# ---------------------------------------------------------------------------
# Cell model
# ---------------------------------------------------------------------------

@dataclass
class Cell:
    sheet: str
    row: int
    col: int
    kind: str              # formula | array | datatable | number | text | bool | date | error_const
    formula: str | None    # A1 formula text incl. leading '=' (None for constants)
    value: Any             # constant value, or cached value for formulas
    array_ref: str | None = None
    r1c1: str | None = None

    @property
    def addr(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def full(self) -> str:
        return f"'{self.sheet}'!{self.addr}"

    @property
    def is_formula(self) -> bool:
        return self.kind in ("formula", "array", "datatable")


def _classify_constant(v) -> str:
    import datetime as _dt
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "number"
    if isinstance(v, (_dt.date, _dt.datetime, _dt.time)):
        return "date"
    if isinstance(v, str) and v.strip() in EXCEL_ERRORS:
        return "error_const"
    return "text"


def iter_cells(wb_f, wb_v, sheets: Iterable[str] | None = None):
    """Yield Cell objects for every non-empty cell."""
    names = list(sheets) if sheets else wb_f.sheetnames
    for name in names:
        ws_f = wb_f[name]
        ws_v = wb_v[name] if name in wb_v.sheetnames else None
        vcells = getattr(ws_v, "_cells", {}) if ws_v is not None else {}
        for (r, c), cell in sorted(ws_f._cells.items()):
            v = cell.value
            if v is None:
                continue
            cached = vcells[(r, c)].value if (r, c) in vcells else None
            if isinstance(v, ArrayFormula):
                text = v.text if v.text.startswith("=") else "=" + v.text
                yield Cell(name, r, c, "array", text, cached, array_ref=v.ref)
            elif DataTableFormula and isinstance(v, DataTableFormula):
                yield Cell(name, r, c, "datatable", "{=TABLE()}", cached)
            elif isinstance(v, str) and v.startswith("=") and len(v) > 1:
                yield Cell(name, r, c, "formula", v, cached)
            else:
                yield Cell(name, r, c, _classify_constant(v), None, v)


def load_cells(path: str, password: str | None = None, sheets=None):
    wb_f, wb_v, real = load_books(path, password)
    cells = list(iter_cells(wb_f, wb_v, sheets))
    for c in cells:
        if c.is_formula and c.kind != "datatable":
            c.r1c1 = to_r1c1(c.formula, c.row, c.col)
    return wb_f, wb_v, real, cells


def index_cells(cells):
    """{sheet: {(row, col): Cell}}"""
    idx: dict[str, dict[tuple[int, int], Cell]] = defaultdict(dict)
    for c in cells:
        idx[c.sheet][(c.row, c.col)] = c
    return idx


# ---------------------------------------------------------------------------
# Reference parsing and R1C1 normalisation
# ---------------------------------------------------------------------------

_CELL = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_COLS = re.compile(r"^(\$?)([A-Za-z]{1,3}):(\$?)([A-Za-z]{1,3})$")
_ROWS = re.compile(r"^(\$?)(\d+):(\$?)(\d+)$")
_R1C1_CELL = re.compile(r"R(\[-?\d+\]|\d+)C(\[-?\d+\]|\d+)")


def split_sheet(operand: str) -> tuple[str | None, str | None, str]:
    """Split an operand into (external_book, sheet, address).

    Handles 'Sheet name'!A1, Sheet!A1, [1]Sheet!A1, '[Book.xlsx]Sheet'!A1,
    'C:\\path\\[Book.xlsx]Sheet'!A1.
    """
    if "!" not in operand:
        return None, None, operand
    prefix, addr = operand.rsplit("!", 1)
    if prefix.startswith("'") and prefix.endswith("'"):
        prefix = prefix[1:-1].replace("''", "'")
    book = None
    m = re.match(r"^(.*)\[([^\]]+)\](.*)$", prefix)
    if m:
        book = (m.group(1) + "[" + m.group(2) + "]") if m.group(1) else m.group(2)
        prefix = m.group(3)
    return book, prefix, addr


def parse_address(addr: str):
    """Return ('cell', (r,c,rabs,cabs)) / ('area', ((r1,c1,..),(r2,c2,..))) /
    ('cols', (c1,c2,abs1,abs2)) / ('rows', (r1,r2,abs1,abs2)) / ('other', addr)."""
    if ":" in addr:
        a, b = addr.split(":", 1)
        ma, mb = _CELL.match(a), _CELL.match(b)
        if ma and mb:
            return "area", (_cell_tuple(ma), _cell_tuple(mb))
        mc = _COLS.match(addr)
        if mc:
            return "cols", (column_index_from_string(mc.group(2).upper()),
                            column_index_from_string(mc.group(4).upper()),
                            bool(mc.group(1)), bool(mc.group(3)))
        mr = _ROWS.match(addr)
        if mr:
            return "rows", (int(mr.group(2)), int(mr.group(4)), bool(mr.group(1)), bool(mr.group(3)))
        return "other", addr
    m = _CELL.match(addr)
    if m:
        return "cell", _cell_tuple(m)
    return "other", addr


def _cell_tuple(m):
    return (int(m.group(4)), column_index_from_string(m.group(2).upper()),
            bool(m.group(3)), bool(m.group(1)))  # row, col, row_abs, col_abs


def _r1c1_part(r, c, rabs, cabs, row, col):
    rp = f"R{r}" if rabs else f"R[{r - row}]"
    cp = f"C{c}" if cabs else f"C[{c - col}]"
    return rp + cp


def operand_to_r1c1(operand: str, row: int, col: int) -> str:
    book, sheet, addr = split_sheet(operand)
    kind, data = parse_address(addr)
    if kind == "cell":
        new = _r1c1_part(*data, row, col)
    elif kind == "area":
        (r1, c1, ra1, ca1), (r2, c2, ra2, ca2) = data
        new = _r1c1_part(r1, c1, ra1, ca1, row, col) + ":" + _r1c1_part(r2, c2, ra2, ca2, row, col)
    elif kind == "cols":
        c1, c2, a1_, a2_ = data
        new = (f"C{c1}" if a1_ else f"C[{c1 - col}]") + ":" + (f"C{c2}" if a2_ else f"C[{c2 - col}]")
    elif kind == "rows":
        r1, r2, a1_, a2_ = data
        new = (f"R{r1}" if a1_ else f"R[{r1 - row}]") + ":" + (f"R{r2}" if a2_ else f"R[{r2 - row}]")
    else:
        return operand
    pre = ""
    if sheet is not None:
        pre = (f"[{book}]" if book else "") + f"'{sheet}'!"
    return pre + new


def tokenize(formula: str):
    try:
        return Tokenizer(formula).items
    except Exception:
        return None


def to_r1c1(formula: str, row: int, col: int) -> str:
    toks = tokenize(formula)
    if toks is None:
        return formula
    out = []
    for t in toks:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            out.append(operand_to_r1c1(t.value, row, col))
        elif t.type == Token.WSPACE:
            continue
        else:
            out.append(t.value)
    return "=" + "".join(out)


def r1c1_to_a1(r1c1: str, row: int, col: int) -> str:
    """Convert a relative-R1C1 normalised formula back to A1 for a target cell."""
    def cell_sub(m):
        rs, cs = m.group(1), m.group(2)
        if rs.startswith("["):
            r = row + int(rs[1:-1]); rabs = ""
        else:
            r = int(rs); rabs = "$"
        if cs.startswith("["):
            c = col + int(cs[1:-1]); cabs = ""
        else:
            c = int(cs); cabs = "$"
        if r < 1 or c < 1:
            return "#REF!"
        return f"{cabs}{get_column_letter(c)}{rabs}{r}"
    return _R1C1_CELL.sub(cell_sub, r1c1)


def refs_in_formula(formula: str):
    """Yield (operand_text, book, sheet, kind, data) for each range operand."""
    toks = tokenize(formula) or []
    for t in toks:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            book, sheet, addr = split_sheet(t.value)
            kind, data = parse_address(addr)
            yield t.value, book, sheet, kind, data


def functions_in_formula(formula: str) -> list[str]:
    toks = tokenize(formula) or []
    return [t.value[:-1].upper() for t in toks
            if t.type == Token.FUNC and t.subtype == Token.OPEN and t.value.endswith("(")]


# ---------------------------------------------------------------------------
# Function-call parser (for lookups / aggregations)
# ---------------------------------------------------------------------------

@dataclass
class Call:
    name: str
    args: list[str] = field(default_factory=list)
    arg_tokens: list[list] = field(default_factory=list)
    depth: int = 0


def parse_calls(formula: str) -> list[Call]:
    """Return every function call in the formula with its top-level argument texts."""
    toks = tokenize(formula)
    if toks is None:
        return []
    calls: list[Call] = []
    stack: list[tuple[Call | None, list]] = []
    for t in toks:
        if t.type == Token.FUNC and t.subtype == Token.OPEN:
            call = Call(t.value[:-1].upper(), depth=len(stack))
            for _, cur in stack:
                cur.append(t)
            stack.append((call, []))
            calls.append(call)
            continue
        if t.type == Token.PAREN and t.subtype == Token.OPEN:
            for _, cur in stack:
                cur.append(t)
            stack.append((None, []))
            continue
        if (t.type == Token.FUNC or t.type == Token.PAREN) and t.subtype == Token.CLOSE:
            if stack:
                call, cur = stack.pop()
                if call is not None:
                    call.arg_tokens.append(cur)
                    call.args.append("".join(x.value for x in cur if x.type != Token.WSPACE))
            for _, cur in stack:
                cur.append(t)
            continue
        if t.type == Token.SEP and t.subtype == Token.ARG and stack and stack[-1][0] is not None:
            call, cur = stack[-1]
            call.arg_tokens.append(cur)
            call.args.append("".join(x.value for x in cur if x.type != Token.WSPACE))
            stack[-1] = (call, [])
            for _, outer in stack[:-1]:
                outer.append(t)
            continue
        for _, cur in stack:
            cur.append(t)
    for c in calls:
        if c.args == [""]:
            c.args, c.arg_tokens = [], []
    return calls


# ---------------------------------------------------------------------------
# Workbook geometry helpers
# ---------------------------------------------------------------------------

def area_bounds(kind, data, ws_max_row=1048576, ws_max_col=16384):
    """Return (r1, c1, r2, c2) for a parsed reference."""
    if kind == "cell":
        r, c = data[0], data[1]
        return r, c, r, c
    if kind == "area":
        (r1, c1, *_), (r2, c2, *_) = data
        return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)
    if kind == "cols":
        return 1, data[0], ws_max_row, data[1]
    if kind == "rows":
        return data[0], 1, data[1], ws_max_col
    return None


def a1(r, c):
    return f"{get_column_letter(c)}{r}"


def header_for(idx_sheet: dict, row: int, col: int, max_up: int = 40) -> str:
    """Best-effort label for a column: nearest text cell above `row` in `col`."""
    for r in range(row - 1, max(0, row - max_up), -1):
        c = idx_sheet.get((r, col))
        if c is not None and c.kind == "text" and str(c.value).strip():
            return str(c.value).strip()
    return ""


def row_label(idx_sheet: dict, row: int, max_cols: int = 6) -> str:
    """Best-effort label for a row: first text cells in the leftmost columns."""
    parts = []
    for col in range(1, max_cols + 1):
        c = idx_sheet.get((row, col))
        if c is not None and c.kind == "text" and str(c.value).strip():
            parts.append(str(c.value).strip())
    return " | ".join(parts)


# ---------------------------------------------------------------------------
# Output helpers
# ---------------------------------------------------------------------------

def ensure_dir(p: str) -> str:
    os.makedirs(p, exist_ok=True)
    return p


def write_json(path: str, obj) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, indent=2, default=str, ensure_ascii=False)


def write_csv(path: str, rows: list[dict], fields: list[str] | None = None) -> None:
    import csv
    if not rows:
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("(no rows)\n")
        return
    fields = fields or list({k: None for r in rows for k in r}.keys())
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in fields})


def short(s, n=160):
    s = "" if s is None else str(s)
    return s if len(s) <= n else s[: n - 1] + "…"


def read_zip_text(path: str, member: str) -> str | None:
    try:
        with zipfile.ZipFile(path) as z:
            return z.read(member).decode("utf-8", errors="replace")
    except KeyError:
        return None


def zip_members(path: str) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return z.namelist()


def not_excel(application) -> bool:
    """True when docProps/app.xml says the file was written by something other than Excel
    (e.g. "Microsoft Excel Compatible / Openpyxl 3.1.5")."""
    a = (application or "").lower()
    return (not a.startswith("microsoft excel")) or any(x in a for x in ("openpyxl", "compatible", "xlsxwriter", "libreoffice", "pandas"))
