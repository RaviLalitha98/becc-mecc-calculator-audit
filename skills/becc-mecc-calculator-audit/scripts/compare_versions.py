#!/usr/bin/env python3
"""Workbook version comparison (stages 9, 10, 25, 26).

Compares two calculator versions (or two "twin" sheets of one workbook, e.g.
Super structure vs Sub structure) at three levels:

  1. SYNTAX     — is the A1 formula text different?
  2. STRUCTURE  — after aligning inserted/deleted rows and columns, does the old
                  formula, with references translated to the new layout, read the
                  same cells? (shifted-but-equivalent formulas drop out)
  3. SEMANTICS  — left to the auditor; the script surfaces the cases that need it:
                  function changes (VLOOKUP→XLOOKUP), reference changes, constant
                  (emission-factor) changes, formulas overwritten by values, with
                  the row/column labels of each cell.

Workbook-level differences: saving application (non-Excel writers such as
openpyxl silently drop features), sheets added/removed/renamed, visibility,
protection, data-validation counts, hidden rows/cols, charts, dynamic-array
metadata, cached values stripped, calc settings, date system, defined names,
external links, newly used version-gated functions, VBA/Power Query.

Outputs: version_diff.md, version_diff_cells.csv
Usage:
  python compare_versions.py OLD.xlsx NEW.xlsx [--out DIR] [--sheets "A,B"]
  python compare_versions.py BOOK.xlsx BOOK.xlsx --pair "Super structure=Sub structure"
"""
from __future__ import annotations

import argparse
import difflib
import os
import re
from collections import Counter, defaultdict

from common import (not_excel, load_cells, index_cells, tokenize, Token, split_sheet, a1, functions_in_formula,
                    header_for, row_label, write_csv, ensure_dir, short)
from excel_compat import base_name, FIRST_VERSION
from inventory import workbook_xml_info, sheet_xml_info, external_links, package_features
from openpyxl.utils import get_column_letter, column_index_from_string


def row_keys(sidx, max_row):
    rows = defaultdict(list)
    for (r, c), cell in sidx.items():
        rows[r].append(cell)
    keys = {}
    for r in range(1, max_row + 1):
        cs = sorted(rows.get(r, []), key=lambda x: x.col)
        labels = [str(x.value).strip()[:40] for x in cs if x.kind == "text" and x.col <= 8 and str(x.value).strip()]
        sigs = [x.r1c1[:60] for x in cs if x.is_formula and x.r1c1][:6]
        keys[r] = "L:" + "|".join(labels[:4]) + "#F:" + "|".join(sigs)
    return keys


def col_keys(sidx, max_col, header_rows=40):
    keys = {}
    for c in range(1, max_col + 1):
        texts = [str(sidx[(r, c)].value).strip()[:40] for r in range(1, header_rows + 1)
                 if (r, c) in sidx and sidx[(r, c)].kind == "text" and str(sidx[(r, c)].value).strip()]
        keys[c] = "|".join(texts[:3]) or f"#{c}"
    return keys


def align(old_keys, new_keys):
    oi, ni = sorted(old_keys), sorted(new_keys)
    sm = difflib.SequenceMatcher(a=[old_keys[k] for k in oi], b=[new_keys[k] for k in ni], autojunk=False)
    m = {}
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal" or (tag == "replace" and (i2 - i1) == (j2 - j1)):
            for d in range(i2 - i1):
                m[oi[i1 + d]] = ni[j1 + d]
    return m, sm.ratio()


def translate(formula, host_sheet, maps, sheet_map):
    toks = tokenize(formula)
    if toks is None:
        return formula
    out = []
    for t in toks:
        if t.type == Token.WSPACE:
            continue
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            book, sh, addr = split_sheet(t.value)
            if book:
                out.append(t.value); continue
            osheet = sh or host_sheet
            nsheet = sheet_map.get(osheet, osheet)
            rm, cm = maps.get(osheet, ({}, {}))

            def cell(x):
                m = re.match(r"^\$?([A-Za-z]{1,3})\$?(\d+)$", x)
                if not m:
                    return x
                r = int(m.group(2)); c = column_index_from_string(m.group(1).upper())
                nr = rm.get(r) if rm else r
                nc = cm.get(c) if cm else c
                if nr is None or nc is None:
                    return "#GONE"
                return f"{get_column_letter(nc)}{nr}"
            new_addr = ":".join(cell(p) for p in addr.split(":"))
            out.append((f"'{nsheet}'!" if sh is not None else "") + new_addr)
        else:
            out.append(t.value)
    return "=" + "".join(out)


def norm(f):
    if f is None:
        return None
    toks = tokenize(f)
    if toks is None:
        return f.replace("$", "").replace(" ", "")
    out = []
    for t in toks:
        if t.type == Token.WSPACE:
            continue
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            book, sh, addr = split_sheet(v)
            v = (f"'{sh}'!" if sh else "") + addr.replace("$", "")
        out.append(v.upper())
    return "=" + "".join(out)


def text_set(sidx):
    return {str(c.value).strip().lower() for c in sidx.values() if c.kind == "text" and str(c.value).strip()}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("old"); ap.add_argument("new")
    ap.add_argument("--old-password"); ap.add_argument("--new-password")
    ap.add_argument("--pair", action="append", default=[], help="explicit sheet pair OLD=NEW (repeatable)")
    ap.add_argument("--sheets", help="limit cell comparison to these OLD sheet names")
    ap.add_argument("--out", default="audit_out")
    a = ap.parse_args()
    out = ensure_dir(a.out)
    wb_o, wv_o, real_o, cells_o = load_cells(a.old, a.old_password)
    same = os.path.abspath(a.new) == os.path.abspath(a.old)
    wb_n, wv_n, real_n, cells_n = (wb_o, wv_o, real_o, cells_o) if same else load_cells(a.new, a.new_password)
    io_, in_ = index_cells(cells_o), index_cells(cells_n)
    wo, wn = workbook_xml_info(real_o), workbook_xml_info(real_n)
    L = [f"# Version comparison\n\n- OLD: `{os.path.basename(a.old)}`\n- NEW: `{os.path.basename(a.new)}`\n"]

    pairs = {}
    for p in a.pair:
        o, n = p.split("=", 1); pairs[o.strip()] = n.strip()
    olds, news = wb_o.sheetnames, wb_n.sheetnames
    if not a.pair:
        L.append("## Workbook-level changes\n")
        L.append(f"- Saved by: `{wo['application']} {wo['app_version'] or ''}` → `{wn['application']} {wn['app_version'] or ''}`"
                 f" (modified {wo['modified']} → {wn['modified']})")
        if not_excel(wn["application"]):
            L.append("  - ⚠ **NEW was not saved by Excel.** Libraries such as openpyxl silently drop cross-sheet (x14) "
                     "data validations, some charts, dynamic-array metadata and cached values. Treat losses below as "
                     "save damage unless documented; ask for an Excel-saved version.")
        for s in olds:
            if s in news:
                pairs[s] = s
        un_n = [s for s in news if s not in pairs.values()]
        for s in [s for s in olds if s not in pairs]:
            ts = text_set(io_.get(s, {}))
            best, score = None, 0
            for t in un_n:
                tt = text_set(in_.get(t, {}))
                j = len(ts & tt) / max(1, len(ts | tt))
                if j > score:
                    best, score = t, j
            if best and score >= 0.5:
                pairs[s] = best; un_n.remove(best)
                L.append(f"- Probable RENAME: `{s}` → `{best}` (label similarity {score:.2f})")
        L.append(f"- Sheets removed: {[s for s in olds if s not in pairs] or 'none'}")
        L.append(f"- Sheets added: {[s for s in news if s not in pairs.values()] or 'none'}")
        so = {s['name']: s for s in wo['sheets']}; sn = {s['name']: s for s in wn['sheets']}
        dv_o = dv_n = 0
        for o, n in pairs.items():
            if so[o]["state"] != sn[n]["state"]:
                L.append(f"- Visibility `{o}`: {so[o]['state']} → {sn[n]['state']}")
            po, pn = sheet_xml_info(real_o, so[o]["xml"]), sheet_xml_info(real_n, sn[n]["xml"])
            dv_o += len(po["validations"]); dv_n += len(pn["validations"])
            if bool(po["protection"]) != bool(pn["protection"]):
                L.append(f"- Protection element `{o}`: {bool(po['protection'])} → {bool(pn['protection'])} "
                         f"(old flags {po['protection']}; check whether it was actually enforced)")
            if len(po["validations"]) != len(pn["validations"]):
                lost = [v for v in po["validations"] if v["ext"]]
                L.append(f"- Data-validation rules `{o}`: {len(po['validations'])} → {len(pn['validations'])}"
                         + (f" ({len(lost)} of the old rules were x14 cross-sheet lists)" if lost else ""))
            if len(po["hidden_rows"]) != len(pn["hidden_rows"]) or po["hidden_cols"] != pn["hidden_cols"]:
                L.append(f"- Hidden rows/col-groups `{o}`: {len(po['hidden_rows'])}/{len(po['hidden_cols'])} → "
                         f"{len(pn['hidden_rows'])}/{len(pn['hidden_cols'])}")
            if po["cond_formats"] != pn["cond_formats"]:
                L.append(f"- Conditional formats `{o}`: {po['cond_formats']} → {pn['cond_formats']}")
            if po["dynamic_cells"] != pn["dynamic_cells"]:
                L.append(f"- Dynamic-array (spill) cells `{o}`: {po['dynamic_cells']} → {pn['dynamic_cells']} — "
                         "formulas may now behave as single-cell/legacy arrays in Excel")
        L.append(f"- Data-validation rules total: {dv_o} → {dv_n}")
        cached_o = sum(1 for c in cells_o if c.is_formula and c.value is not None)
        cached_n = sum(1 for c in cells_n if c.is_formula and c.value is not None)
        if cached_n < 0.5 * cached_o:
            L.append(f"- Cached formula results: {cached_o} → {cached_n} — NEW was not recalculated/saved by Excel after editing")
        for k in ("calcMode", "fullCalcOnLoad", "iterate", "iterateCount", "iterateDelta", "fullPrecision"):
            if wo["calcPr"].get(k) != wn["calcPr"].get(k):
                L.append(f"- calcPr.{k}: {wo['calcPr'].get(k)} → {wn['calcPr'].get(k)}")
        if wo["date_system"] != wn["date_system"]:
            L.append(f"- DATE SYSTEM changed {wo['date_system']} → {wn['date_system']} (dates shift by 1,462 days)")
        q = lambda s: (s or "").replace("'", "")
        do = {(d['name'], d['localSheetId']): d['refers_to'] for d in wo['defined_names']}
        dn = {(d['name'], d['localSheetId']): d['refers_to'] for d in wn['defined_names']}
        for k in sorted(set(do) | set(dn), key=str):
            if k[0].startswith("_xlchart") and (k in do) and (k in dn):
                continue
            if q(do.get(k)) != q(dn.get(k)):
                L.append(f"- Name `{k[0]}` (scope {k[1]}): `{do.get(k)}` → `{dn.get(k)}`")
        eo = {e['target'] for e in external_links(real_o)}; en = {e['target'] for e in external_links(real_n)}
        if eo != en:
            L.append(f"- External link targets: removed {sorted(eo - en)}, added {sorted(en - eo)}")
        fo = Counter(base_name(f) for c in cells_o if c.is_formula for f in functions_in_formula(c.formula))
        fn_ = Counter(base_name(f) for c in cells_n if c.is_formula for f in functions_in_formula(c.formula))
        newf = sorted(set(fn_) - set(fo))
        L.append(f"- Functions newly used: {newf or 'none'}"
                 + (f" (version-gated: {[f + '≥' + FIRST_VERSION[f] for f in newf if f in FIRST_VERSION]})" if any(f in FIRST_VERSION for f in newf) else ""))
        L.append(f"- Functions no longer used: {sorted(set(fo) - set(fn_)) or 'none'}")
        po_, pn_ = package_features(real_o), package_features(real_n)
        for k in ("vba_project", "power_query", "connections", "charts", "chartex", "drawings", "metadata_xml", "comments"):
            if po_[k] != pn_[k]:
                L.append(f"- {k}: {po_[k]} → {pn_[k]}")
    else:
        L.append("- Explicit sheet pairs: " + ", ".join(f"`{o}` ↔ `{n}`" for o, n in pairs.items()))
    cmp_pairs = {o: n for o, n in pairs.items() if not a.sheets or o in {s.strip() for s in a.sheets.split(",")}}

    maps, align_info = {}, {}
    all_pairs = dict(pairs)
    for o, n in all_pairs.items():
        so_, sn_ = io_.get(o, {}), in_.get(n, {})
        mro = max([r for r, c in so_] or [1]); mrn = max([r for r, c in sn_] or [1])
        mco = max([c for r, c in so_] or [1]); mcn = max([c for r, c in sn_] or [1])
        rmap, rr = align(row_keys(so_, mro), row_keys(sn_, mrn))
        cmap, cr = align(col_keys(so_, mco), col_keys(sn_, mcn))
        if a.pair and rr < 0.3:
            rmap = {r: r for r in range(1, mro + 1)}
        if a.pair and cr < 0.3:
            cmap = {c: c for c in range(1, mco + 1)}
        maps[o] = (rmap, cmap)
        align_info[o] = {"row_ratio": rr, "col_ratio": cr,
                         "ins_rows": sorted(set(range(1, mrn + 1)) - set(rmap.values())),
                         "del_rows": sorted(set(range(1, mro + 1)) - set(rmap)),
                         "ins_cols": sorted(set(range(1, mcn + 1)) - set(cmap.values())),
                         "del_cols": sorted(set(range(1, mco + 1)) - set(cmap)),
                         "row_shift_mode": Counter(v - k for k, v in rmap.items()).most_common(3)}

    rows = []
    for o, n in cmp_pairs.items():
        so_, sn_ = io_.get(o, {}), in_.get(n, {})
        rmap, cmap = maps[o]
        seen_new = set()
        for (r, c), co in so_.items():
            base = {"sheet_old": o, "cell_old": co.addr, "sheet_new": n,
                    "label": short(row_label(so_, r), 50), "header": short(header_for(so_, r, c), 40), "sig": co.r1c1 or ""}
            nr, nc = rmap.get(r), cmap.get(c)
            if nr is None or nc is None:
                if co.is_formula or co.kind == "number":
                    rows.append(dict(base, cell_new="", change="REMOVED", old=short(co.formula or co.value, 200), new="",
                                     note="row/col deleted or unaligned"))
                continue
            cn = sn_.get((nr, nc)); seen_new.add((nr, nc))
            rec = dict(base, cell_new=a1(nr, nc))
            if cn is None:
                if co.is_formula or co.kind == "number":
                    rows.append(dict(rec, change="CLEARED", old=short(co.formula or co.value, 200), new="", note=""))
                continue
            if co.is_formula and cn.is_formula:
                if norm(co.formula) == norm(cn.formula):
                    continue
                tr = translate(co.formula, o, maps, all_pairs)
                if norm(tr) == norm(cn.formula):
                    continue
                fo_ = [base_name(f) for f in functions_in_formula(co.formula)]
                fn2 = [base_name(f) for f in functions_in_formula(cn.formula)]
                if Counter(fo_) != Counter(fn2):
                    kind = "FUNCTION_CHANGE"
                    note = f"functions {sorted(set(fo_) - set(fn2))} → {sorted(set(fn2) - set(fo_))}"
                    if {"VLOOKUP", "HLOOKUP", "INDEX", "MATCH"} & set(fo_) and {"XLOOKUP", "XMATCH"} & set(fn2):
                        note += " — LOOKUP IMPLEMENTATION CHANGE: verify match mode, not-found handling, duplicates"
                else:
                    kind = "REFERENCE_CHANGE"
                    note = "same functions, different references after layout alignment"
                    if "#GONE" in tr:
                        note += " (old formula referenced deleted rows/cols)"
                rows.append(dict(rec, change=kind, old=short(co.formula, 250), new=short(cn.formula, 250),
                                 note=note + f" | old translated: {short(tr, 200)}"))
            elif co.is_formula:
                rows.append(dict(rec, change="FORMULA_TO_VALUE", old=short(co.formula, 250), new=short(cn.value, 60),
                                 note="formula replaced by a constant — possible overwrite"))
            elif cn.is_formula:
                rows.append(dict(rec, change="VALUE_TO_FORMULA", old=short(co.value, 60), new=short(cn.formula, 250), note=""))
            elif co.value != cn.value:
                if isinstance(co.value, (int, float)) and isinstance(cn.value, (int, float)) and not isinstance(co.value, bool):
                    if abs(cn.value - co.value) <= 1e-9 * max(1.0, abs(co.value)):
                        continue  # float representation noise
                    rel = (cn.value - co.value) / co.value * 100 if co.value else float("inf")
                    rows.append(dict(rec, change="DATA_CHANGE", old=co.value, new=cn.value, note=f"{rel:+.2f}%"))
                elif co.kind == "text" and cn.kind == "text":
                    if str(co.value).strip() != str(cn.value).strip():
                        rows.append(dict(rec, change="LABEL_CHANGE", old=short(co.value, 80), new=short(cn.value, 80), note=""))
                else:
                    rows.append(dict(rec, change="TYPE_CHANGE", old=short(co.value, 60), new=short(cn.value, 60), note=""))
        for (r, c), cn in sn_.items():
            if (r, c) not in seen_new and (cn.is_formula or cn.kind == "number"):
                rows.append({"sheet_old": o, "cell_old": "", "sheet_new": n, "cell_new": cn.addr, "change": "ADDED",
                             "old": "", "new": short(cn.formula or cn.value, 200), "note": "",
                             "label": short(row_label(sn_, r), 50), "header": short(header_for(sn_, r, c), 40),
                             "sig": cn.r1c1 or ""})

    write_csv(os.path.join(out, "version_diff_cells.csv"), rows,
              ["sheet_old", "cell_old", "sheet_new", "cell_new", "change", "old", "new", "note", "label", "header"])
    L.append("\n## Layout alignment per sheet\n")
    for o, inf in align_info.items():
        if o in cmp_pairs and (inf["ins_rows"] or inf["del_rows"] or inf["ins_cols"] or inf["del_cols"] or inf["row_ratio"] < 0.98):
            L.append(f"- `{o}`: row-key similarity {inf['row_ratio']:.2f}, col {inf['col_ratio']:.2f}; "
                     f"inserted rows {short(inf['ins_rows'], 80)}, deleted rows {short(inf['del_rows'], 80)}, "
                     f"inserted cols {[get_column_letter(x) for x in inf['ins_cols']][:20]}, "
                     f"deleted cols {[get_column_letter(x) for x in inf['del_cols']][:20]}; dominant row shifts {inf['row_shift_mode']}")
    L.append("\nAlignment is heuristic: when similarity is low, verify a few mappings by hand before trusting REFERENCE_CHANGE.")
    cnt = Counter((r["sheet_old"], r["change"]) for r in rows)
    L.append("\n## Cell-level changes (after alignment; shifted-but-equivalent formulas excluded)\n")
    L.append("| Sheet | Change | Cells |\n|---|---|---|")
    for (s, k), v in sorted(cnt.items()):
        L.append(f"| {s} | {k} | {v} |")
    order = ["FORMULA_TO_VALUE", "FUNCTION_CHANGE", "REFERENCE_CHANGE", "CLEARED", "REMOVED", "DATA_CHANGE",
             "VALUE_TO_FORMULA", "TYPE_CHANGE", "ADDED", "LABEL_CHANGE"]
    grouped = {}
    for r in rows:
        gk = r["sig"] if r["change"] in ("FUNCTION_CHANGE", "REFERENCE_CHANGE", "FORMULA_TO_VALUE", "REMOVED", "ADDED", "CLEARED") else id(r)
        grouped.setdefault((r["sheet_old"], r["change"], gk), []).append(r)
    for kind in order:
        gs = [g for k, g in grouped.items() if k[1] == kind]
        if not gs:
            continue
        L.append(f"\n### {kind} ({sum(len(g) for g in gs)} cells, {len(gs)} groups)\n")
        for g in gs[:80]:
            r = g[0]
            cl = ", ".join((x["cell_old"] or x["cell_new"]) for x in g[:8]) + (f" … (+{len(g) - 8})" if len(g) > 8 else "")
            L.append(f"- **{r['sheet_old']}** [{cl}] {r['label']} / {r['header']}  \n  old: `{r['old']}`  \n  new: `{r['new']}`  \n  {r['note']}")
        if len(gs) > 80:
            L.append(f"… {len(gs) - 80} more groups in version_diff_cells.csv")
    with open(os.path.join(out, "version_diff.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")
    print(f"version diff → {out}/version_diff.md ; " + ", ".join(f"{k}:{v}" for k, v in Counter(r['change'] for r in rows).items()))


if __name__ == "__main__":
    main()
